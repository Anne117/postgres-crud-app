#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Сборка отчёта ДЗ №4 (.docx) из model/report-data.json.

Запуск:  python tools/make_report.py
Результат: hw4/report/hw4-otchet.docx
"""
from __future__ import annotations

import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
HW4 = os.path.dirname(HERE)
MODEL = os.path.join(HW4, "model")
OUT_DIR = os.path.join(HW4, "report")

sys.path.insert(0, HW4)
import catalog as C  # noqa: E402

FONT = "Times New Roman"


def risk_band(score: float) -> str:
    """Качественная шкала риска по сумме DREAD (5..25)."""
    if score >= 20:
        return "Критический"
    if score >= 15:
        return "Высокий"
    if score >= 10:
        return "Средний"
    return "Низкий"


def set_base_font(doc: Document) -> None:
    for name in ("Normal",):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.element.rPr.rFonts.set(qn("w:cs"), FONT)
    for i in range(1, 5):
        try:
            st = doc.styles["Heading %d" % i]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    st = doc.styles["Title"]
    st.font.name = FONT


def para(doc, text, *, style=None, size=None, bold=None, italic=None,
         align=None, space_after=6):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def table(doc, header, rows, *, widths=None, size=9, style="Table Grid",
          align_center_cols=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(size)
        run.font.name = FONT
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            val = "" if val is None else str(val)
            cells[i].text = ""
            first = True
            for line in val.split("\n"):
                p = cells[i].paragraphs[0] if first else cells[i].add_paragraph()
                first = False
                run = p.add_run(line)
                run.font.size = Pt(size)
                run.font.name = FONT
                p.paragraph_format.space_after = Pt(0)
                if i in align_center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Mm(w)
    return t


def landscape(doc: Document) -> None:
    s = doc.sections[-1]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    for name in ("left_margin", "right_margin"):
        setattr(s, name, Mm(15))
    s.top_margin = s.bottom_margin = Mm(15)


def portrait(doc: Document) -> None:
    doc.add_section()
    s = doc.sections[-1]
    s.orientation = WD_ORIENT.PORTRAIT
    if s.page_width > s.page_height:
        s.page_width, s.page_height = s.page_height, s.page_width
    s.left_margin = s.right_margin = Mm(20)
    s.top_margin = s.bottom_margin = Mm(15)


def main() -> int:
    data = json.load(open(os.path.join(MODEL, "report-data.json"),
                          encoding="utf-8"))
    threats = data["threats"]
    top5 = data["top5"]
    dl = data["dread_labels"]
    sr = data["stride_ru"]
    dread_keys = ["D", "R", "E", "A", "D2"]
    os.makedirs(OUT_DIR, exist_ok=True)

    doc = Document()
    doc.sections[0].left_margin = doc.sections[0].right_margin = Mm(20)
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Mm(15)
    set_base_font(doc)

    # --- титул ---------------------------------------------------------
    para(doc, "Отчёт по заданию №4", size=18, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "Моделирование угроз информационной безопасности "
              "по методологии STRIDE и оценка рисков по методике DREAD",
         size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    para(doc, "Объект моделирования: приложение «Склад интернет-магазина» "
              "(CRUD-приложение к PostgreSQL), разработанное в рамках "
              "задания №3.\n"
              "Инструмент: OWASP Threat Dragon (формат модели v2), "
              "методология STRIDE.\n"
              "Источник сведений об угрозах: Банк данных угроз безопасности "
              "информации ФСТЭК России (https://bdu.fstec.ru/threat).",
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    para(doc, "Состав приложенных исходных файлов модели: "
              "dfd-crud-app.drawio (диаграмма потоков данных), "
              "threat-dragon-crud-app.json (модель OWASP Threat Dragon), "
              "dfd-crud-app.svg / dfd-crud-app.png (экспорт диаграммы).",
         size=10, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

    # --- 1. объект моделирования --------------------------------------
    doc.add_heading("1. Краткое описание объекта моделирования", level=1)
    para(doc, "Моделируемое приложение — консольное CRUD-приложение к СУБД "
              "PostgreSQL, развёрнутое в Docker (docker compose). "
              "Состав компонентов:")
    for item in [
        "crud_app — контейнер приложения (Python, модули app.py — меню и "
        "сценарии, db.py — подключение к PostgreSQL, sqlsafe.py — "
        "безопасная сборка SQL, logger.py — журналирование);",
        "postgres — контейнер СУБД PostgreSQL с базой данных приложения "
        "(таблицы supplier, product, \"order\", order_item);",
        "crud.env — файл конфигурации приложения "
        "(строка подключения, учётные данные БД);",
        "app.docker.conf — конфигурационный файл СУБД;",
        "том ./logs — каталог технических журналов приложения.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    para(doc, "Роли доступа к СУБД: пользователь приложения (ограниченные "
              "права) и служебная административная роль app_admin. "
              "Взаимодействие пользователь—приложение—СУБД образует "
              "цепочку потоков данных, промоделированную на диаграмме "
              "раздела 3.")

    # --- 2. методология -----------------------------------------------
    doc.add_heading("2. Методика моделирования и инструмент", level=1)
    para(doc, "Диаграмма потоков данных (DFD) построена в нотации, "
              "поддерживаемой OWASP Threat Dragon: внешние сущности "
              "(actor), процессы (process), хранилища данных (store), "
              "потоки данных (flow) и границы доверия (trust boundary).")
    para(doc, "Методология STRIDE применена покомпонентно: для каждого "
              "элемента DFD проанализированы угрозы шести категорий "
              "методологии.")
    table(doc,
          ["Обозначение", "Категория STRIDE", "Нарушаемое свойство ИБ"],
          [["S", sr["S"], "Аутентичность"],
           ["T", sr["T"], "Целостность"],
           ["R", sr["R"], "Неотказуемость (доказуемость)"],
           ["I", sr["I"], "Конфиденциальность"],
           ["D", sr["D"], "Доступность"],
           ["E", sr["E"], "Авторизация (права доступа)"]],
          widths=[25, 85, 55], size=10, align_center_cols=(0,))
    para(doc, "Инструмент OWASP Threat Dragon формирует для элементов "
              "диаграммы типовые (шаблонные) угрозы STRIDE. Полученные "
              "шаблонные угрозы локализованы: переведены на русский язык "
              "и уточнены применительно к архитектуре именно этого "
              "приложения (указаны конкретные элементы, файлы, таблицы и "
              "роли). Локализованный перечень приведён в разделе 4; те же "
              "локализованные угрозы сохранены в исходном файле модели "
              "threat-dragon-crud-app.json (свойство threats у элементов "
              "диаграммы).")

    # --- 3. диаграмма --------------------------------------------------
    doc.add_heading("3. Диаграмма потоков данных", level=1)
    para(doc, "На рисунке 1 представлена диаграмма потоков данных "
              "приложения. Обозначения: прямоугольник — внешняя сущность "
              "(actor), овал/скруглённый блок — процесс, фигура со двойной "
              "вертикальной чертой — хранилище данных (store), "
              "пунктирная рамка — граница доверия, сплошная стрелка — "
              "поток данных, красная пунктирная стрелка — поток "
              "воздействия нарушителя (пересекает границу доверия). "
              "Внутри элементов в квадратных скобках указаны "
              "идентификаторы связанных угроз (T-NN).")
    png = os.path.join(MODEL, "dfd-crud-app.png")
    if os.path.exists(png):
        doc.add_picture(png, width=Mm(235))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Рисунок 1 — Диаграмма потоков данных CRUD-приложения "
              "(методология STRIDE)", size=10, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_heading("3.1. Элементы диаграммы", level=1)
    kind_ru = {"actor": "Внешняя сущность",
               "process": "Процесс",
               "store": "Хранилище данных",
               "text": "Пояснение"}
    table(doc,
          ["ID", "Тип", "Элемент", "Назначение"],
          [[e["id"], kind_ru.get(e["kind"], e["kind"]),
            e["label"].replace("\n", " ").replace("[", "").replace("]", ""),
            e["desc"]] for e in data["elements"]],
          widths=[26, 30, 45, 109], size=9, align_center_cols=(0, 1))

    doc.add_heading("3.2. Потоки данных", level=1)
    name_of = {e["id"]: e["label"].split("\n")[0] for e in data["elements"]}
    table(doc,
          ["№", "Источник", "Приёмник", "Поток данных",
           "Пересекает границу доверия"],
          [[f["n"], name_of.get(f["src"], f["src"]),
            name_of.get(f["dst"], f["dst"]),
            f["label"].split(". ", 1)[-1].replace("\n", " "),
            "да" if f["crosses"] else "нет"] for f in data["flows"]],
          widths=[10, 40, 40, 90, 30], size=9, align_center_cols=(0, 4))

    # --- 4. локализованные угрозы STRIDE -------------------------------
    portrait(doc)
    doc.add_heading("4. Локализованные угрозы по методологии STRIDE",
                    level=1)
    para(doc, "Ниже приведён полный локализованный перечень угроз, "
              "полученных на основе диаграммы (шаблонные угрозы Threat "
              "Dragon уточнены под конкретную архитектуру). В столбце "
              "«Категория» указано обозначение STRIDE, в столбце "
              "«Элемент» — идентификатор элемента диаграммы, к которому "
              "отнесена угроза.")
    table(doc,
          ["ID", "Кат.", "Элемент", "Угроза", "Описание"],
          [[t["id"], t["stride"], t["target"], t["title"], t["desc"]]
           for t in threats],
          widths=[11, 11, 24, 48, 76], size=8.5,
          align_center_cols=(0, 1))
    para(doc, "Меры противодействия по каждой угрозе приведены в "
              "приложении А.", size=10, italic=True)

    # --- 5. DREAD ------------------------------------------------------
    doc.add_heading("5. Оценка рисков по методике DREAD", level=1)
    para(doc, "Каждая угроза оценена по пяти критериям методики DREAD. "
              "Принята целочисленная шкала от 1 до 5, где 1 — минимальное, "
              "5 — максимальное значение критерия. Общая оценка риска — "
              "сумма баллов по пяти критериям (от 5 до 25).")
    table(doc,
          ["Критерий", "Смысл критерия", "Шкала"],
          [["D — Damage", dl["D"],
            "1 — незначительный ущерб; 5 — полная потеря данных/системы"],
           ["R — Reproducibility", dl["R"],
            "1 — воспроизводится редко; 5 — воспроизводится всегда"],
           ["E — Exploitability", dl["E"],
            "1 — требуется высокая квалификация и особые условия; "
            "5 — эксплуатируется простыми средствами"],
           ["A — Affected users", dl["A"],
            "1 — единичный пользователь; 5 — все пользователи и данные"],
           ["D — Discoverability", dl["D2"],
            "1 — обнаружить крайне сложно; 5 — уязвимость очевидна"]],
          widths=[42, 78, 45], size=9.5)
    para(doc, "Качественная шкала итогового риска: 20–25 — «Критический», "
              "15–19 — «Высокий», 10–14 — «Средний», 5–9 — «Низкий».",
         size=10, italic=True)

    doc.add_heading("5.1. Таблица расчёта рисков (оценки критериев DREAD)",
                    level=1)
    para(doc, "Таблица содержит оценочные значения всех пяти критериев "
              "для каждой угрозы и итоговую сумму баллов. Строки "
              "отсортированы по убыванию общей оценки риска.")
    rows = []
    for i, t in enumerate(threats, 1):
        d = dict(zip(dread_keys, t["dread"]))
        rows.append([i, t["id"], t["stride"], t["title"],
                     d["D"], d["R"], d["E"], d["A"], d["D2"],
                     t["score"], risk_band(t["score"])])
    table(doc,
          ["№", "ID", "Кат.", "Угроза", "D", "R", "E", "A", "D",
           "Сумма", "Риск"],
          rows, widths=[8, 11, 11, 60, 8, 8, 8, 8, 8, 12, 20], size=8.5,
          align_center_cols=(0, 1, 2, 4, 5, 6, 7, 8, 9, 10))
    para(doc, "Столбцы D, R, E, A, D — оценки критериев Damage, "
              "Reproducibility, Exploitability, Affected users и "
              "Discoverability соответственно.", size=9, italic=True)

    # --- 6. итоговая таблица -------------------------------------------
    doc.add_heading("6. Итоговая таблица угроз с общей оценкой риска",
                    level=1)
    para(doc, "Итоговая таблица угроз, отсортированная по убыванию общей "
              "оценки риска.")
    rows = []
    for i, t in enumerate(threats, 1):
        rows.append([i, t["id"], sr[t["stride"]].split(" — ")[0],
                     t["title"], t["score"], risk_band(t["score"])])
    table(doc,
          ["Место", "ID", "Категория STRIDE", "Угроза",
           "Общая оценка (5–25)", "Уровень риска"],
          rows, widths=[14, 11, 48, 62, 20, 25], size=9,
          align_center_cols=(0, 1, 4, 5))

    # --- 7. БДУ ФСТЭК ---------------------------------------------------
    doc.add_heading("7. Соответствие пяти самых критичных угроз банку данных "
                    "угроз ФСТЭК России", level=1)
    para(doc, "Для пяти угроз с наибольшей общей оценкой риска найдены "
              "соответствующие угрозы безопасности информации (УБИ) в "
              "Банке данных угроз безопасности информации ФСТЭК России "
              "(bdu.fstec.ru). Соответствие установлено по способу "
              "воздействия нарушителя и по объекту воздействия.")
    table(doc,
          ["Место", "ID", "Угроза (локализованная)", "Риск",
           "УБИ БДУ ФСТЭК", "Наименование УБИ", "Обоснование соответствия",
           "Ссылка"],
          [[p["place"], p["id"], p["title"], p["score"],
            p["bdu"]["ubi"], p["bdu"]["title"], p["bdu"]["why"],
            p["bdu"]["url"]] for p in top5],
          widths=[10, 11, 42, 11, 18, 42, 62, 34], size=8,
          align_center_cols=(0, 1, 3, 4))
    para(doc, "Полное соответствие для всех 15 угроз приведено в "
              "приложении Б.", size=10, italic=True)

    # --- 8. выводы ------------------------------------------------------
    doc.add_heading("8. Выводы", level=1)
    crit = [t for t in threats if risk_band(t["score"]) == "Критический"]
    high = [t for t in threats if risk_band(t["score"]) == "Высокий"]
    para(doc, "Всего на диаграмме локализовано %d угроз STRIDE. "
              "К категории «Критический» отнесены %d угроз(и), "
              "к категории «Высокий» — %d, остальные — «Средний». "
              "Наибольший риск связан с угрозами нарушения целостности "
              "данных СУБД (SQL-инъекция, несанкционированная "
              "модификация данных), подмены конфигурации и повышения "
              "привилегий (использование административной роли, побег из "
              "контейнера). Все пять наиболее критичных угроз имеют "
              "прямые соответствия в БДУ ФСТЭК России."
              % (len(threats), len(crit), len(high)))
    para(doc, "Часть угроз уже закрыта мерами, реализованными в "
              "приложении (белый список колонок и операторов, "
              "параметризация значений через psycopg sql.Literal, "
              "ограничение прав роли пользователя, журналирование "
              "операций). Рекомендуемые остаточные меры приведены по "
              "каждой угрозе в приложении А.")

    # --- приложения -----------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Приложение А. Карточки угроз: описание, "
                    "оценка DREAD и меры противодействия", level=1)
    for i, t in enumerate(threats, 1):
        d = dict(zip(dread_keys, t["dread"]))
        h = doc.add_heading("А.%d. %s — %s" % (i, t["id"], t["title"]),
                            level=2)
        h.paragraph_format.space_before = Pt(10)
        table(doc,
              ["Параметр", "Значение"],
              [["Категория STRIDE", "%s — %s" % (t["stride"],
                                                 sr[t["stride"]])],
               ["Элемент DFD", t["target"]],
               ["Описание угрозы", t["desc"]],
               ["Оценки DREAD (D, R, E, A, D)",
                "%d, %d, %d, %d, %d" % tuple(d[k] for k in dread_keys)],
               ["Общая оценка риска",
                "%s (уровень: %s)" % (t["score"], risk_band(t["score"]))],
               ["Меры противодействия", t["mitigation"]],
               ["Соответствие УБИ БДУ ФСТЭК",
                ("%s «%s» (%s)" % (t["bdu"]["ubi"], t["bdu"]["title"],
                                   t["bdu"]["url"])) if t.get("bdu")
                else "прямого соответствия не подобрано"]],
              widths=[48, 117], size=9)

    doc.add_page_break()
    doc.add_heading("Приложение Б. Соответствие всех угроз УБИ БДУ ФСТЭК "
                    "России", level=1)
    table(doc,
          ["ID", "Угроза", "Общая оценка", "УБИ БДУ ФСТЭК",
           "Наименование УБИ"],
          [[t["id"], t["title"], t["score"],
            (t["bdu"] or {}).get("ubi", "—"),
            (t["bdu"] or {}).get("title", "—")] for t in threats],
          widths=[11, 62, 18, 20, 54], size=9,
          align_center_cols=(0, 2, 3))

    doc.add_page_break()
    doc.add_heading("Приложение В. Состав приложенных исходных файлов "
                    "модели", level=1)
    table(doc,
          ["Файл", "Назначение"],
          [["model/dfd-crud-app.drawio",
            "Исходная диаграмма потоков данных (редактируется в "
            "draw.io / diagrams.net)"],
           ["model/threat-dragon-crud-app.json",
            "Исходный файл модели OWASP Threat Dragon v2 (диаграмма + "
            "локализованные угрозы STRIDE на элементах)"],
           ["model/dfd-crud-app.svg",
            "Векторный экспорт диаграммы"],
           ["model/dfd-crud-app.png",
            "Растровый экспорт диаграммы (использован в отчёте)"],
           ["model/report-data.json",
            "Машиночитаемые данные отчёта (элементы, потоки, угрозы, "
            "оценки DREAD, соответствия УБИ)"]],
          widths=[62, 103], size=9)

    out = os.path.join(OUT_DIR, "hw4-otchet.docx")
    doc.save(out)
    print("docx: %s (%d байт)" % (out, os.path.getsize(out)))
    print("таблиц: %d, абзацев: %d" % (len(doc.tables), len(doc.paragraphs)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
